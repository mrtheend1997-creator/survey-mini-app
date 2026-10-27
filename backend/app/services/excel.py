"""Excel export of survey results (summary + one row per response)."""

from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..config import get_settings
from ..models import Answer, Response, Survey, SurveyVersion, User
from .answers import readable
from .results import merged_questions, region_names, survey_results

L = {
    "uz": {
        "summary": "Umumiy",
        "responses": "Javoblar",
        "survey": "So'rovnoma",
        "participants": "Ishtirokchilar",
        "completed": "Yakunlangan",
        "in_progress": "Yakunlanmagan",
        "rate": "Yakunlash foizi",
        "question": "Savol",
        "option": "Javob varianti",
        "count": "Soni",
        "percent": "Foiz",
        "average": "O'rtacha",
        "no": "№",
        "status": "Holati",
        "started": "Boshlangan",
        "finished": "Yakunlangan vaqti",
        "version": "Versiya",
        "respondent": "Ishtirokchi",
        "phone": "Telefon",
        "st_completed": "Yakunlangan",
        "st_in_progress": "Jarayonda",
        "st_timed_out": "Vaqt tugagan",
        "anonymous": "Anonim",
        "row": "Qator",
        "avg_place": "O'rtacha o'rin",
        "nps": "NPS indeksi",
    },
    "ru": {
        "summary": "Сводка",
        "responses": "Ответы",
        "survey": "Опрос",
        "participants": "Участники",
        "completed": "Завершено",
        "in_progress": "Не завершено",
        "rate": "Процент завершения",
        "question": "Вопрос",
        "option": "Вариант ответа",
        "count": "Количество",
        "percent": "Процент",
        "average": "Среднее",
        "no": "№",
        "status": "Статус",
        "started": "Начато",
        "finished": "Завершено (время)",
        "version": "Версия",
        "respondent": "Участник",
        "phone": "Телефон",
        "st_completed": "Завершена",
        "st_in_progress": "В процессе",
        "st_timed_out": "Время истекло",
        "anonymous": "Анонимно",
        "row": "Строка",
        "avg_place": "Среднее место",
        "nps": "Индекс NPS",
    },
}

HEADER_FILL = PatternFill("solid", fgColor="0B2A4A")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _header(ws, row: int, values: list[str]) -> None:
    for col, v in enumerate(values, 1):
        c = ws.cell(row=row, column=col, value=v)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")


def build_workbook(db: Session, survey: Survey, lang: str = "ru", version_id: int | None = None) -> bytes:
    t = L[lang]
    tz = ZoneInfo(get_settings().app_timezone)
    res = survey_results(db, survey, version_id)
    questions = merged_questions(db, survey, version_id)
    title = getattr(survey, f"title_{lang}") or survey.title_uz

    wb = Workbook()
    ws = wb.active
    ws.title = t["summary"]
    ws["A1"] = t["survey"]
    ws["B1"] = title
    ws["A1"].font = Font(bold=True)
    rows = [
        (t["participants"], res["participants"]),
        (t["completed"], res["completed"]),
        (t["in_progress"], res["in_progress"]),
        (t["rate"], f"{res['completion_rate']}%"),
    ]
    for i, (k, v) in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=2, value=v)
    r = len(rows) + 3
    _header(ws, r, [t["question"], t["option"], t["count"], t["percent"]])
    r += 1
    for idx, q in enumerate(res["questions"], 1):
        qtext = f"{idx}. {q.get(f'text_{lang}') or q['text_uz']}"
        stats = q["stats"]
        if stats["kind"] == "none":
            continue
        if stats["kind"] == "matrix":
            for row in stats["rows"]:
                rtext = row.get(f"text_{lang}") or row["text_uz"]
                for col, cell in zip(stats["columns"], row["cells"]):
                    ctext = col.get(f"text_{lang}") or col["text_uz"]
                    ws.append([qtext, f"{rtext} — {ctext}", cell["count"], cell["percent"] / 100])
                    ws.cell(row=ws.max_row, column=4).number_format = "0.0%"
                    qtext = ""
            continue
        if stats["kind"] in ("ranking", "points"):
            for item in stats["items"]:
                extra = item.get("average") if stats["kind"] == "ranking" and not stats.get("pairwise") else item["percent"] / 100
                ws.append([qtext, item.get(f"text_{lang}") or item["text_uz"], item["count"], extra])
                if stats["kind"] != "ranking" or stats.get("pairwise"):
                    ws.cell(row=ws.max_row, column=4).number_format = "0.0%"
                qtext = ""
            continue
        if stats["kind"] in ("choice", "rating"):
            for item in stats["items"]:
                ws.append([qtext, item.get(f"text_{lang}") or item["text_uz"], item["count"], item["percent"] / 100])
                ws.cell(row=ws.max_row, column=4).number_format = "0.0%"
                qtext = ""
            if stats.get("average") is not None:
                ws.append(["", t["average"], stats["average"], None])
            if stats.get("nps") and stats["nps"]["score"] is not None:
                ws.append(["", t["nps"], stats["nps"]["score"], None])
        elif stats["kind"] == "number":
            ws.append([qtext, t["average"], stats["average"], None])
        else:
            ws.append([qtext, "", q["answered"], None])
        r = ws.max_row + 1
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 12

    ws2 = wb.create_sheet(t["responses"])
    anonymous = survey.survey_type == "anonymous"
    head = [t["no"], t["status"], t["started"], t["finished"], t["version"]]
    if not anonymous:
        head += [t["respondent"], "Telegram", t["phone"]]
    questions = [q for q in questions if q["qtype"] != "info"]
    head += [f"{i}. {q.get(f'text_{lang}') or q['text_uz']}" for i, q in enumerate(questions, 1)]
    regions = region_names(db) if any(q["qtype"] == "region" for q in questions) else {}
    _header(ws2, 1, head)

    rq = select(Response).where(Response.survey_id == survey.id)
    if version_id:
        rq = rq.where(Response.version_id == version_id)
    responses = list(db.scalars(rq.order_by(Response.id)))
    versions = {v.id: v.number for v in db.scalars(select(SurveyVersion).where(SurveyVersion.survey_id == survey.id))}
    answers: dict[int, dict[int, dict]] = {}
    if responses:
        for a in db.scalars(select(Answer).where(Answer.response_id.in_([x.id for x in responses]))):
            answers.setdefault(a.response_id, {})[a.question_id] = a.value
    users = {}
    if not anonymous:
        uids = {x.user_id for x in responses if x.user_id}
        if uids:
            users = {u.id: u for u in db.scalars(select(User).where(User.id.in_(uids)))}

    for n, resp in enumerate(responses, 1):
        row = [
            n,
            t.get(f"st_{resp.status}", t["st_in_progress"]),
            resp.started_at.astimezone(tz).strftime("%Y-%m-%d %H:%M") if resp.started_at else "",
            resp.completed_at.astimezone(tz).strftime("%Y-%m-%d %H:%M") if resp.completed_at else "",
            versions.get(resp.version_id, ""),
        ]
        if not anonymous:
            u = users.get(resp.user_id)
            info = resp.respondent_info or {}
            row += [
                info.get("full_name") or (u.full_name if u else ""),
                f"@{u.username}" if u and u.username else (str(u.telegram_id) if u else ""),
                info.get("phone", ""),
            ]
        ra = answers.get(resp.id, {})
        row += [readable(q, ra.get(q["id"]), lang, regions) for q in questions]
        ws2.append(row)
    for col in range(1, len(head) + 1):
        ws2.column_dimensions[get_column_letter(col)].width = 18 if col <= 8 else 32
    ws2.freeze_panes = "B2"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
