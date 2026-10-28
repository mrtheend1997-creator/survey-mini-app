"""Jumps between questions, sections and the bundled e-commerce questionnaire."""

from io import BytesIO

from openpyxl import load_workbook

from app.db import SessionLocal
from app.models import Answer, AppSetting, Response
from app.services import importer, logic

from .conftest import SUPER_ID, auth

ADMIN = auth(SUPER_ID, "Admin")
USER = auth(7301, "Resp")


def _code(r):
    return r.json()["detail"]["code"]


def seed(client):
    client.get("/api/me", headers=ADMIN)
    with SessionLocal() as db:
        ids = importer.seed_surveys(db)
        assert importer.seed_surveys(db) == []  # only once
        assert db.get(AppSetting, "seed:elektron_tijorat") is not None
    assert len(ids) == 1
    sid = ids[0]
    s = client.get(f"/api/admin/surveys/{sid}", headers=ADMIN).json()
    return s


def items(s):
    """key -> question, numbering questions 1..20 and sections s1..s5 like the Word file."""
    out, n, sec = {}, 0, 0
    for q in s["questions"]:
        if q["qtype"] == "section":
            sec += 1
            out[f"s{sec}"] = q
        else:
            n += 1
            out[n] = q
    return out


def test_seeded_questionnaire_matches_word_file(client):
    s = seed(client)
    assert s["status"] == "draft" and s["survey_type"] == "anonymous" and not s["allow_repeat"]
    assert s["title_uz"].startswith("Oʻzbekiston Respublikasida elektron tijoratning")
    assert s["description_uz"].startswith("Hurmatli respondent!") and "Izoh:" in s["description_uz"]
    assert s["thanks_uz"] == "Ishtirokingiz uchun rahmat!"
    it = items(s)
    assert len([k for k in it if isinstance(k, int)]) == 20 and len([k for k in it if isinstance(k, str)]) == 5
    assert it["s3"]["description_uz"].startswith("Ushbu boʻlimga faqat")
    assert it[3]["qtype"] == "short_text"
    multi = {k for k in range(1, 21) if it[k]["qtype"] == "multiple"}
    assert multi == {10, 11, 14, 15, 16, 17, 20}
    assert {k for k in range(1, 21) if it[k]["max_choices"] == 3} == {14, 17, 20}
    assert [k for k in range(1, 21) if it[k]["allow_other"]] == [10]
    assert [o["text_uz"] for o in it[8]["options"]] == ["Ha", "Yoʻq, lekin ilgari bunday xaridlarni amalga oshirganman", "Hech qachon amalga oshirmaganman"]
    assert it[14]["hint_uz"] == "Koʻpi bilan uchta javobni tanlang."
    # jumps point at the sections IV and V
    o8 = [o["id"] for o in it[8]["options"]]
    assert it[8]["logic"] == {
        "rules": [{"if": {"op": "is", "options": [o8[1]]}, "to": it["s4"]["id"]}, {"if": {"op": "is", "options": [o8[2]]}, "to": it["s4"]["id"]}],
        "default": None,
    }
    assert it[15]["logic"] == {"rules": [], "default": it["s5"]["id"]}
    check = client.get(f"/api/admin/surveys/{s['id']}/logic-check", headers=ADMIN).json()
    assert check == {"errors": [], "warnings": []}


def walk(client, sid, choice8, user=USER):
    """Answer like a respondent following the app's route; returns the numbers of the questions shown."""
    s = client.get(f"/api/admin/surveys/{sid}", headers=ADMIN).json()
    it = items(s)
    number = {q["id"]: k for k, q in it.items() if isinstance(k, int)}
    resp = client.post(f"/api/surveys/{sid}/responses", json={"language": "uz"}, headers=user).json()
    token = resp["token"]
    put = lambda q, v: client.put(f"/api/responses/{token}/answers/{q['id']}", json={"value": v}, headers=user)
    # the same route the app follows, worked out from the snapshot's jumps
    shown, i, qs = [], 0, resp["questions"]
    index = {q["id"]: n for n, q in enumerate(qs)}
    while i < len(qs):
        q = qs[i]
        target = None
        if q["qtype"] != "section":
            shown.append(number[q["id"]])
            k = number[q["id"]]
            if q["qtype"] == "short_text":
                v = {"text": "Toshkent shahri"}
            elif q["qtype"] == "single":
                v = {"option_id": q["options"][choice8 if k == 8 else 0]["id"]}
            else:
                v = {"option_ids": [o["id"] for o in q["options"][:2]]}
            assert put(q, v).status_code == 200, (k, put(q, v).text)
            target = logic.next_target(q, v)
        i = index[target] if target else i + 1
    return token, shown, it


def test_both_routes(client):
    sid = seed(client)["id"]
    assert client.post(f"/api/admin/surveys/{sid}/publish", headers=ADMIN).status_code == 200
    token, shown, it = walk(client, sid, 0)
    assert shown == [1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 18, 19, 20]
    assert client.post(f"/api/responses/{token}/submit", headers=USER).json()["ok"]

    for n, choice in ((2, 1), (3, 2)):
        user = auth(7310 + n)
        token, shown, _ = walk(client, sid, choice, user)
        assert shown == [1, 2, 3, 4, 5, 6, 7, 8, 16, 17, 18, 19, 20]
        r = client.post(f"/api/responses/{token}/submit", headers=user)
        assert r.status_code == 200 and r.json()["ok"], r.text

    with SessionLocal() as db:
        a = db.query(Response).filter_by(token=token).one()
        assert sorted(a.skipped) == sorted(it[k]["id"] for k in range(9, 16))
        text = db.query(Answer).filter_by(response_id=a.id, question_id=it[3]["id"]).one()
        assert text.value == {"text": "Toshkent shahri"}

    results = client.get(f"/api/admin/surveys/{sid}/results", headers=ADMIN).json()
    by_id = {q["id"]: q for q in results["questions"]}
    assert by_id[it[9]["id"]]["answered"] == 1 and by_id[it[9]["id"]]["not_applicable"] == 2
    assert by_id[it[16]["id"]]["answered"] == 2 and by_id[it[16]["id"]]["not_applicable"] == 1
    assert by_id[it[18]["id"]]["answered"] == 3

    # anonymous: no names or Telegram ids anywhere for the admins
    rows = client.get(f"/api/admin/surveys/{sid}/responses", headers=ADMIN).json()
    assert rows["anonymous"] and all(r["respondent"] is None for r in rows["items"])
    detail = client.get(f"/api/admin/responses/{rows['items'][0]['id']}?lang=ru", headers=ADMIN).json()
    assert detail["respondent"] is None
    statuses = {i["question_id"]: i["status"] for i in detail["items"]}
    assert statuses[it[9]["id"]] == "not_applicable" and statuses[it[16]["id"]] == "answered"
    wb = load_workbook(BytesIO(client.get(f"/api/admin/surveys/{sid}/export.xlsx?lang=uz", headers=ADMIN).content))
    ws = wb.worksheets[1]
    header = [c.value for c in ws[1]]
    assert "Telegram" not in header
    flat = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert "Tegishli emas" in flat
    assert not any(x in v for v in flat for x in ("7301", "7312", "7313", "Resp"))


def test_changing_the_branch_drops_old_answers(client):
    sid = seed(client)["id"]
    client.post(f"/api/admin/surveys/{sid}/publish", headers=ADMIN)
    token, shown, it = walk(client, sid, 0)  # answered 9-15 on the "Ha" branch
    # back to question 8: now "never bought online", then 16 and 17
    o8 = [o["id"] for o in it[8]["options"]]
    client.put(f"/api/responses/{token}/answers/{it[8]['id']}", json={"value": {"option_id": o8[2]}}, headers=USER)
    r = client.post(f"/api/responses/{token}/submit", headers=USER)
    assert r.status_code == 422 and {m["question_id"] for m in r.json()["detail"]["missing"]} == {it[16]["id"], it[17]["id"]}
    for k in (16, 17):
        client.put(f"/api/responses/{token}/answers/{it[k]['id']}", json={"value": {"option_ids": [it[k]["options"][0]["id"]]}}, headers=USER)
    assert client.post(f"/api/responses/{token}/submit", headers=USER).json()["ok"]
    with SessionLocal() as db:
        resp = db.query(Response).filter_by(token=token).one()
        kept = {a.question_id for a in db.query(Answer).filter_by(response_id=resp.id)}
    assert not kept & {it[k]["id"] for k in range(9, 16)}
    assert {it[k]["id"] for k in (1, 8, 16, 17, 18, 20)} <= kept


def test_max_three_answers(client):
    sid = seed(client)["id"]
    client.post(f"/api/admin/surveys/{sid}/publish", headers=ADMIN)
    it = items(client.get(f"/api/admin/surveys/{sid}", headers=ADMIN).json())
    token = client.post(f"/api/surveys/{sid}/responses", json={"language": "uz"}, headers=USER).json()["token"]
    for k in (14, 17, 20):
        four = [o["id"] for o in it[k]["options"][:4]]
        r = client.put(f"/api/responses/{token}/answers/{it[k]['id']}", json={"value": {"option_ids": four}}, headers=USER)
        assert r.status_code == 422, k
        r = client.put(f"/api/responses/{token}/answers/{it[k]['id']}", json={"value": {"option_ids": four[:3]}}, headers=USER)
        assert r.status_code == 200, k
    # no limit where the Word file has none
    r = client.put(f"/api/responses/{token}/answers/{it[15]['id']}", json={"value": {"option_ids": [o["id"] for o in it[15]["options"][:5]]}}, headers=USER)
    assert r.status_code == 200


def test_editing_jumps_in_the_builder(client):
    s = seed(client)
    sid = s["id"]
    it = items(s)

    def body(q, **changes):
        b = {k: q[k] for k in ("text_uz", "text_ru", "qtype", "required", "hint_uz", "hint_ru", "min_choices", "max_choices", "allow_other", "settings", "logic")}
        b["options"] = [{"id": o["id"], "text_uz": o["text_uz"], "text_ru": o["text_ru"]} for o in q["options"]]
        return {**b, **changes}

    # a backward jump would loop: refused at publish and reported by the check
    r = client.put(f"/api/admin/questions/{it[12]['id']}", json=body(it[12], logic={"rules": [], "default": it[9]["id"]}), headers=ADMIN)
    assert r.status_code == 200
    check = client.get(f"/api/admin/surveys/{sid}/logic-check", headers=ADMIN).json()
    assert check["errors"] == [{"question_id": it[12]["id"], "code": "logic_backward", "target_id": it[9]["id"]}]
    r = client.post(f"/api/admin/surveys/{sid}/publish", headers=ADMIN)
    assert r.status_code == 422 and _code(r) == "logic_backward"
    # a jump on an answer of a free-text question makes no sense
    r = client.put(f"/api/admin/questions/{it[3]['id']}", json=body(it[3], logic={"rules": [{"if": {"op": "is", "options": [5]}, "to": "end"}]}), headers=ADMIN)
    assert r.status_code == 422 and _code(r) == "invalid_logic"
    client.put(f"/api/admin/questions/{it[12]['id']}", json=body(it[12], logic=None), headers=ADMIN)

    # questions 16-17 become unreachable if nobody is sent there
    r = client.put(f"/api/admin/questions/{it[8]['id']}", json=body(it[8], logic=None), headers=ADMIN)
    assert r.status_code == 200
    warn = client.get(f"/api/admin/surveys/{sid}/logic-check", headers=ADMIN).json()["warnings"]
    assert {w["question_id"] for w in warn} == {it["s4"]["id"], it[16]["id"], it[17]["id"]}

    # an admin without logic_manage may edit the text but not the jumps
    client.get("/api/me", headers=auth(7320))
    perms = dict(client.get("/api/admin/admins", headers=ADMIN).json()["templates"]["create_edit"], logic_manage=False)
    a = client.post("/api/admin/admins", json={"telegram_id": 7320, "role": "admin", "permissions": perms, "access": "all", "survey_level": "edit"}, headers=ADMIN)
    assert a.status_code == 200, a.text
    ed = auth(7320)
    q8 = client.get(f"/api/admin/surveys/{sid}", headers=ed)
    if q8.status_code != 200:  # this admin isn't given the survey: the Super Admin shares it
        client.patch(f"/api/admin/surveys/{sid}", json={"manager_ids": [a.json()["id"]]}, headers=ADMIN)
    it = items(client.get(f"/api/admin/surveys/{sid}", headers=ed).json())
    r = client.put(f"/api/admin/questions/{it[8]['id']}", json=body(it[8], text_uz=it[8]["text_uz"] + " "), headers=ed)
    assert r.status_code == 200, r.text
    o8 = [o["id"] for o in it[8]["options"]]
    r = client.put(f"/api/admin/questions/{it[8]['id']}", json=body(it[8], logic={"jumps": [{"when": o8[1], "to": it["s4"]["id"]}]}), headers=ed)
    assert r.status_code == 403


def test_duplicate_keeps_jumps(client):
    s = seed(client)
    copy = client.post(f"/api/admin/surveys/{s['id']}/duplicate", headers=ADMIN).json()
    c = client.get(f"/api/admin/surveys/{copy['id']}", headers=ADMIN).json()
    it = items(c)
    o8 = [o["id"] for o in it[8]["options"]]
    assert it[8]["logic"]["rules"][0] == {"if": {"op": "is", "options": [o8[1]]}, "to": it["s4"]["id"]}
    assert client.get(f"/api/admin/surveys/{copy['id']}/logic-check", headers=ADMIN).json()["errors"] == []
    # a copied question gets its own option ids in its jumps
    dup = client.post(f"/api/admin/questions/{it[8]['id']}/duplicate", headers=ADMIN).json()
    assert dup["logic"]["rules"][0]["if"]["options"] == [dup["options"][1]["id"]]


def test_import_endpoint(client):
    data = {
        "title": "Test",
        "questions": [
            {"key": "a", "qtype": "yes_no", "text": "Savol?", "jumps": [{"when": "no", "to": "end"}]},
            {"key": "b", "qtype": "short_text", "text": "Nega?"},
        ],
    }
    r = client.post("/api/admin/surveys/import", json=data, headers=ADMIN)
    assert r.status_code == 200, r.text
    s = client.get(f"/api/admin/surveys/{r.json()['id']}", headers=ADMIN).json()
    assert s["questions"][0]["logic"] == {"rules": [{"if": {"op": "is", "values": ["no"]}, "to": "end"}], "default": None}
    bad = client.post("/api/admin/surveys/import", json={"title": "x", "questions": [{"key": "a", "text": "q", "options": ["1"], "jumps": [{"when_option": 0, "to": "zzz"}]}]}, headers=ADMIN)
    assert bad.status_code == 422 and _code(bad) == "invalid_import"
    assert client.post("/api/admin/surveys/import", json=data, headers=USER).status_code == 403

    # "no" ends the survey: the second question is not applicable
    client.post(f"/api/admin/surveys/{s['id']}/publish", headers=ADMIN)
    token = client.post(f"/api/surveys/{s['id']}/responses", json={"language": "ru"}, headers=USER).json()["token"]
    client.put(f"/api/responses/{token}/answers/{s['questions'][0]['id']}", json={"value": {"bool": False}}, headers=USER)
    assert client.post(f"/api/responses/{token}/submit", headers=USER).json()["ok"]
