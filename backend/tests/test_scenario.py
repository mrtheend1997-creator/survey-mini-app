"""The readiness scenario from the specification, executed through the API."""

from io import BytesIO

from openpyxl import load_workbook

from .conftest import SUPER_ID, auth

ADMIN = auth(SUPER_ID, "Admin")


def choice(uz, ru, n=4):
    return {
        "text_uz": uz,
        "text_ru": ru,
        "qtype": "single",
        "required": True,
        "options": [{"text_uz": f"Variant {i}", "text_ru": f"Вариант {i}"} for i in range(1, n + 1)],
    }


def create_survey(client, title="So'rovnoma"):
    r = client.post(
        "/api/admin/surveys",
        json={
            "title_uz": title,
            "title_ru": f"Опрос {title}",
            "description_uz": "Tavsif",
            "description_ru": "Описание",
            "survey_type": "anonymous",
            "allow_repeat": False,
        },
        headers=ADMIN,
    )
    assert r.status_code == 200, r.text
    return r.json()


def test_full_admin_scenario(client):
    # 3-4. create survey with uz/ru title
    s = create_survey(client)
    sid = s["id"]
    assert s["status"] == "draft"

    # 5-6. add 10 questions with options
    qids = []
    for i in range(10):
        r = client.post(f"/api/admin/surveys/{sid}/questions", json=choice(f"Savol {i}", f"Вопрос {i}"), headers=ADMIN)
        assert r.status_code == 200, r.text
        qids.append(r.json()["id"])
    qs = client.get(f"/api/admin/surveys/{sid}/questions", headers=ADMIN).json()
    assert [q["id"] for q in qs] == qids

    # 7. reorder: move last question up, then full reorder
    r = client.post(f"/api/admin/questions/{qids[9]}/move", json={"direction": "up"}, headers=ADMIN)
    assert [q["id"] for q in r.json()][8] == qids[9]
    new_order = list(reversed(qids))
    r = client.post(f"/api/admin/surveys/{sid}/questions/reorder", json={"question_ids": new_order}, headers=ADMIN)
    assert [q["id"] for q in r.json()] == new_order

    # 8-9. edit a question and delete one option
    q0 = next(q for q in qs if q["id"] == qids[0])
    opts = q0["options"][:3]  # drop the 4th option
    opts[0]["text_ru"] = "Отлично"
    body = {**choice("Yangi matn", "Новый текст"), "options": opts}
    r = client.put(f"/api/admin/questions/{qids[0]}", json=body, headers=ADMIN)
    assert r.status_code == 200, r.text
    assert len(r.json()["options"]) == 3
    assert r.json()["options"][0]["text_ru"] == "Отлично"
    assert r.json()["text_ru"] == "Новый текст"

    # duplicate + delete a question
    dup = client.post(f"/api/admin/questions/{qids[1]}/duplicate", headers=ADMIN).json()
    assert client.delete(f"/api/admin/questions/{dup['id']}", headers=ADMIN).status_code == 200

    # survey not visible to users before publishing
    user = auth(5001, "User")
    assert client.get("/api/surveys", headers=user).json() == []

    # 11. publish
    r = client.post(f"/api/admin/surveys/{sid}/publish", headers=ADMIN)
    assert r.status_code == 200, r.text
    assert r.json()["status"] == "active"
    assert r.json()["version_number"] == 1

    # 12. take the survey from another Telegram account
    lst = client.get("/api/surveys", headers=user).json()
    assert len(lst) == 1 and lst[0]["question_count"] == 10
    r = client.post(f"/api/surveys/{sid}/responses", json={"language": "ru"}, headers=user)
    assert r.status_code == 200, r.text
    token = r.json()["token"]
    questions = r.json()["questions"]
    assert [q["id"] for q in questions] == new_order

    # submit too early -> required check
    r = client.post(f"/api/responses/{token}/submit", headers=user)
    assert r.status_code == 422 and r.json()["detail"]["code"] == "required_missing"

    for idx, q in enumerate(questions):
        r = client.put(
            f"/api/responses/{token}/answers/{q['id']}",
            json={"value": {"option_id": q["options"][0]["id"]}, "current_index": idx},
            headers=user,
        )
        assert r.status_code == 200, r.text

    # resume returns saved answers (draft recovery)
    r = client.post(f"/api/surveys/{sid}/responses", json={"token": token}, headers=user)
    assert len(r.json()["answers"]) == 10

    # 13. submit, then again (double submit is harmless)
    assert client.post(f"/api/responses/{token}/submit", headers=user).json()["already"] is False
    assert client.post(f"/api/responses/{token}/submit", headers=user).json()["already"] is True
    # can't answer or start again (single participation)
    r = client.put(f"/api/responses/{token}/answers/{questions[0]['id']}", json={"value": None}, headers=user)
    assert r.status_code == 409
    r = client.post(f"/api/surveys/{sid}/responses", json={}, headers=user)
    assert r.status_code == 409 and r.json()["detail"]["code"] == "already_completed"
    assert client.get("/api/surveys", headers=user).json()[0]["completed"] is True

    # anonymous: no user id stored on the response
    from app.db import SessionLocal
    from app.models import Response

    with SessionLocal() as db:
        assert db.query(Response).one().user_id is None

    # 14-15. results and dashboard
    res = client.get(f"/api/admin/surveys/{sid}/results", headers=ADMIN).json()
    assert res["participants"] == 1 and res["completed"] == 1 and res["completion_rate"] == 100.0
    first = res["questions"][0]["stats"]["items"][0]
    assert first["count"] == 1 and first["percent"] == 100.0
    dash = client.get("/api/admin/dashboard", headers=ADMIN).json()
    assert dash["kpi"]["completed"] == 1 and dash["kpi"]["today"] == 1 and dash["kpi"]["active_surveys"] == 1

    # 16. Excel
    for lang in ("ru", "uz"):
        r = client.get(f"/api/admin/surveys/{sid}/export.xlsx?lang={lang}", headers=ADMIN)
        assert r.status_code == 200
        wb = load_workbook(BytesIO(r.content))
        assert len(wb.worksheets) == 2
        ws = wb.worksheets[1]
        assert ws.max_row == 2  # header + one response

    # responses viewer
    lst = client.get(f"/api/admin/surveys/{sid}/responses", headers=ADMIN).json()
    assert lst["total"] == 1 and lst["anonymous"] is True and lst["items"][0]["respondent"] is None
    detail = client.get(f"/api/admin/responses/{lst['items'][0]['id']}?lang=ru", headers=ADMIN).json()
    assert len(detail["items"]) == 10 and detail["items"][-1]["answer"] == "Отлично"

    # 17-18. second survey managed independently
    s2 = create_survey(client, "Ikkinchi")
    client.post(f"/api/admin/surveys/{s2['id']}/questions", json=choice("A", "Б"), headers=ADMIN)
    client.post(f"/api/admin/surveys/{s2['id']}/publish", headers=ADMIN)
    client.post(f"/api/admin/surveys/{sid}/status", json={"status": "paused"}, headers=ADMIN)
    lst = client.get("/api/surveys", headers=user).json()
    assert [x["id"] for x in lst] == [s2["id"]]

    # audit log was written
    log = client.get("/api/admin/audit", headers=ADMIN).json()
    actions = {i["action"] for i in log["items"]}
    assert {"survey.create", "question.create", "survey.publish", "results.export"} <= actions or log["total"] > 30


def test_versioning_keeps_old_answers(client):
    s = create_survey(client)
    sid = s["id"]
    q = client.post(f"/api/admin/surveys/{sid}/questions", json=choice("Eski", "Старый", 2), headers=ADMIN).json()
    client.post(f"/api/admin/surveys/{sid}/publish", headers=ADMIN)

    u1 = auth(7001)
    r = client.post(f"/api/surveys/{sid}/responses", json={}, headers=u1).json()
    old_option = q["options"][1]["id"]
    client.put(f"/api/responses/{r['token']}/answers/{q['id']}", json={"value": {"option_id": old_option}}, headers=u1)
    client.post(f"/api/responses/{r['token']}/submit", headers=u1)

    # admin edits published survey: removes the option that was chosen, adds a question
    body = {**choice("Yangi", "Новый", 0), "options": [q["options"][0], {"text_uz": "Uch", "text_ru": "Три"}]}
    client.put(f"/api/admin/questions/{q['id']}", json=body, headers=ADMIN)
    client.post(f"/api/admin/surveys/{sid}/questions", json=choice("Qo'shimcha", "Дополнительный", 2), headers=ADMIN)
    info = client.get(f"/api/admin/surveys/{sid}", headers=ADMIN).json()
    assert info["has_unpublished_changes"] is True

    # users still get version 1 until changes are published
    assert client.get("/api/surveys", headers=auth(7002)).json()[0]["question_count"] == 1
    r = client.post(f"/api/admin/surveys/{sid}/publish", headers=ADMIN).json()
    assert r["version_number"] == 2 and r["has_unpublished_changes"] is False
    assert client.get("/api/surveys", headers=auth(7002)).json()[0]["question_count"] == 2

    # old answer still counted, the removed option still has its label
    res = client.get(f"/api/admin/surveys/{sid}/results", headers=ADMIN).json()
    items = res["questions"][0]["stats"]["items"]
    removed = next(i for i in items if i["option_id"] == old_option)
    assert removed["count"] == 1 and removed["removed"] is True
    detail_id = client.get(f"/api/admin/surveys/{sid}/responses", headers=ADMIN).json()["items"][0]["id"]
    detail = client.get(f"/api/admin/responses/{detail_id}?lang=ru", headers=ADMIN).json()
    assert detail["version_number"] == 1 and detail["items"][0]["text_ru"] == "Старый"
    assert detail["items"][0]["answer"] == "Вариант 2"


def test_question_types_validation(client):
    s = create_survey(client)
    sid = s["id"]
    defs = [
        {"text_uz": "a", "text_ru": "a", "qtype": "multiple", "min_choices": 1, "max_choices": 2, "allow_other": True,
         "options": [{"text_uz": "x", "text_ru": "x"}, {"text_uz": "y", "text_ru": "y"}, {"text_uz": "z", "text_ru": "z"}]},
        {"text_uz": "b", "text_ru": "b", "qtype": "text"},
        {"text_uz": "c", "text_ru": "c", "qtype": "number"},
        {"text_uz": "d", "text_ru": "d", "qtype": "yes_no"},
        {"text_uz": "e", "text_ru": "e", "qtype": "rating5"},
        {"text_uz": "f", "text_ru": "f", "qtype": "rating10"},
        {"text_uz": "g", "text_ru": "g", "qtype": "dropdown", "options": [{"text_uz": "1", "text_ru": "1"}, {"text_uz": "2", "text_ru": "2"}]},
        {"text_uz": "h", "text_ru": "h", "qtype": "date", "required": False},
    ]
    for d in defs:
        r = client.post(f"/api/admin/surveys/{sid}/questions", json=d, headers=ADMIN)
        assert r.status_code == 200, r.text
    # invalid: choice with one option, missing translation
    bad = client.post(f"/api/admin/surveys/{sid}/questions", json=choice("x", "y", 1), headers=ADMIN)
    assert bad.json()["detail"]["code"] == "options_required"
    bad = client.post(f"/api/admin/surveys/{sid}/questions", json={**choice("x", "", 2)}, headers=ADMIN)
    assert bad.json()["detail"]["code"] == "question_text_required"

    client.post(f"/api/admin/surveys/{sid}/publish", headers=ADMIN)
    u = auth(8001)
    r = client.post(f"/api/surveys/{sid}/responses", json={}, headers=u).json()
    t, qs = r["token"], r["questions"]

    def put(i, value):
        return client.put(f"/api/responses/{t}/answers/{qs[i]['id']}", json={"value": value}, headers=u)

    o = [x["id"] for x in qs[0]["options"]]
    assert put(0, {"option_ids": o}).status_code == 422  # more than max 2
    assert put(0, {"option_ids": [o[0]], "other": "boshqa"}).status_code == 200
    assert put(1, {"text": "  matn "}).json()["value"] == {"text": "matn"}
    assert put(2, {"number": "abc"}).status_code == 422
    assert put(2, {"number": 12.5}).status_code == 200
    assert put(3, {"bool": True}).status_code == 200
    assert put(4, {"rating": 6}).status_code == 422
    assert put(4, {"rating": 5}).status_code == 200
    assert put(5, {"rating": 10}).status_code == 200
    assert put(6, {"option_id": 999999}).status_code == 422
    assert put(6, {"option_id": qs[6]["options"][1]["id"]}).status_code == 200
    assert put(7, {"date": "not-a-date"}).status_code == 422
    # date is optional
    assert client.post(f"/api/responses/{t}/submit", headers=u).status_code == 200
    res = client.get(f"/api/admin/surveys/{sid}/results", headers=ADMIN).json()
    assert res["questions"][4]["stats"]["average"] == 5
    assert res["questions"][2]["stats"]["average"] == 12.5
