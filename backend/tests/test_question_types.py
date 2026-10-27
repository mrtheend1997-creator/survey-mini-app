"""Extended question types: settings validation, answers, submit rules, results, export, files."""

from io import BytesIO

from openpyxl import load_workbook

from app.db import SessionLocal
from app.models import Region

from .conftest import SUPER_ID, auth

ADMIN = auth(SUPER_ID, "Admin")
USER = auth(7001, "Resp")


def opts(n=3, image=False):
    return [{"text_uz": f"V{i}", "text_ru": f"В{i}", **({"image_url": f"/api/uploads/{i:032x}.png"} if image else {})} for i in range(1, n + 1)]


def q(qtype, settings=None, options=None, required=True, **extra):
    return {"text_uz": f"Savol {qtype}", "text_ru": f"Вопрос {qtype}", "qtype": qtype, "required": required, "settings": settings, "options": options or [], **extra}


QUESTIONS = {
    "single": q("single", options=opts(), allow_other=True),
    "multiple": q("multiple", options=opts(4), min_choices=2, max_choices=3),
    "dropdown": q("dropdown", options=opts()),
    "image_choice": q("image_choice", {"multiple": True}, opts(3, image=True)),
    "short_text": q("short_text", {"format": "email", "max_len": 100}),
    "text": q("text", {"min_len": 10, "max_len": 500}),
    "scale": q("scale", {"min": 0, "max": 10, "label_min_uz": "Yomon", "label_min_ru": "Плохо"}),
    "stars": q("stars", {"count": 7}),
    "nps": q("nps"),
    "yes_no": q("yes_no", {"unsure": True}),
    "date": q("date", {"min_date": "2020-01-01", "max_date": "2030-12-31"}),
    "time": q("time", {"min_time": "08:00", "max_time": "18:00"}),
    "number": q("number", {"min": 1, "max": 100, "integer": True, "unit_uz": "yil", "unit_ru": "лет"}),
    "matrix": q("matrix", {"rows": [{"text_uz": "Xizmat", "text_ru": "Сервис"}, {"text_uz": "Narx", "text_ru": "Цена"}]}, opts(3)),
    "matrix_multi": q("matrix_multi", {"rows": [{"text_uz": "Qator", "text_ru": "Строка"}]}, opts(3)),
    "ranking": q("ranking", options=opts(3)),
    "points": q("points", {"total": 100}, opts(3)),
    "pairwise": q("pairwise", options=opts(3)),
    "file": q("file", {"formats": ["pdf", "png"], "max_mb": 1}, required=False),
    "video": q("video", {"max_len": 200}, required=False, video_url="https://youtu.be/dQw4w9WgXcQ"),
    "contact": q("contact", {"fields": {"full_name": True, "phone": True}, "purpose_uz": "Bog'lanish uchun", "purpose_ru": "Для связи"}),
    "region": q("region", {"district": True}),
    "info": q("info", description_uz="Ma'lumot", description_ru="Информация"),
    "rating5": q("rating5"),
}


def seed_regions():
    with SessionLocal() as db:
        tash = Region(name_uz="Toshkent shahri", name_ru="г. Ташкент", position=0)
        sam = Region(name_uz="Samarqand viloyati", name_ru="Самаркандская обл.", position=1)
        db.add_all([tash, sam])
        db.flush()
        d = Region(parent_id=tash.id, name_uz="Chilonzor", name_ru="Чиланзар", position=0)
        db.add(d)
        db.commit()
        return tash.id, sam.id, d.id


def make_survey(client, survey_type="named"):
    r = client.post(
        "/api/admin/surveys",
        json={"title_uz": "Turlar", "title_ru": "Типы", "survey_type": survey_type, "allow_repeat": True},
        headers=ADMIN,
    )
    assert r.status_code == 200, r.text
    return r.json()["id"]


def test_all_types_end_to_end(client):
    tash, sam, chil = seed_regions()
    sid = make_survey(client)
    ids = {}
    for name, body in QUESTIONS.items():
        r = client.post(f"/api/admin/surveys/{sid}/questions", json=body, headers=ADMIN)
        assert r.status_code == 200, (name, r.text)
        ids[name] = r.json()
    # matrix rows got stable ids, info is never required
    assert [row["id"] for row in ids["matrix"]["settings"]["rows"]] == ["r1", "r2"]
    assert ids["info"]["required"] is False
    assert ids["image_choice"]["options"][0]["image_url"].startswith("/api/uploads/")
    # a hidden draft question is not published
    hidden = client.post(f"/api/admin/surveys/{sid}/questions", json=q("short_text", hidden=True), headers=ADMIN).json()
    assert hidden["hidden"] is True

    r = client.post(f"/api/admin/surveys/{sid}/publish", headers=ADMIN)
    assert r.status_code == 200, r.text

    resp = client.post(f"/api/surveys/{sid}/responses", json={"language": "ru"}, headers=USER).json()
    token = resp["token"]
    snap_ids = [x["id"] for x in resp["questions"]]
    assert hidden["id"] not in snap_ids and len(snap_ids) == len(QUESTIONS)
    regions = client.get("/api/regions", headers=USER).json()
    assert regions[0]["districts"][0]["id"] == chil

    def put(name, value):
        return client.put(f"/api/responses/{token}/answers/{ids[name]['id']}", json={"value": value}, headers=USER)

    o = {name: [x["id"] for x in body["options"]] for name, body in ids.items()}
    rows = [row["id"] for row in ids["matrix"]["settings"]["rows"]]

    # malformed answers are rejected at once
    bad = [
        ("scale", {"rating": 11}),
        ("stars", {"rating": 8}),
        ("nps", {"rating": -1}),
        ("time", {"time": "25:00"}),
        ("time", {"time": "07:00"}),
        ("date", {"date": "2019-05-05"}),
        ("matrix", {"matrix": {"r9": o["matrix"][0]}}),
        ("ranking", {"order": o["ranking"][:2]}),
        ("points", {"points": {str(o["points"][0]): 101}}),
        ("pairwise", {"pairs": {f"{o['pairwise'][0]}-{o['pairwise'][1]}": o["pairwise"][2]}}),
        ("contact", {"contact": {"email": "a@b.c"}}),
        ("region", {"region": chil}),
        ("region", {"region": sam, "district": chil}),
        ("file", {"files": [{"key": "answer/999/x.pdf", "name": "x.pdf", "size": 1}]}),
        ("yes_no", {"unsure": "yes"}),
        ("image_choice", {"option_ids": [999]}),
    ]
    for name, value in bad:
        assert put(name, value).status_code == 422, (name, value)

    # the file question: upload a real PDF, a fake one is refused
    files_url = f"/api/responses/{token}/files/{ids['file']['id']}"
    fake = client.post(files_url, files={"file": ("cv.pdf", b"not a pdf", "application/pdf")}, headers=USER)
    assert fake.status_code == 415
    exe = client.post(files_url, files={"file": ("x.exe", b"MZ\x90\x00", "application/octet-stream")}, headers=USER)
    assert exe.status_code == 415
    big = client.post(files_url, files={"file": ("big.pdf", b"%PDF-1.4" + b"0" * (1024 * 1024 + 1), "application/pdf")}, headers=USER)
    assert big.status_code == 413
    up = client.post(files_url, files={"file": ("cv.pdf", b"%PDF-1.4 test", "application/pdf")}, headers=USER)
    assert up.status_code == 200, up.text
    uploaded = up.json()
    # another respondent can't attach this file
    other = client.post(f"/api/surveys/{sid}/responses", json={"language": "ru"}, headers=auth(7002)).json()
    r = client.put(f"/api/responses/{other['token']}/answers/{ids['file']['id']}", json={"value": {"files": [uploaded]}}, headers=auth(7002))
    assert r.status_code == 422

    good = {
        "single": {"other": "O'zim"},
        "multiple": {"option_ids": o["multiple"][:1]},  # fewer than min_choices: saved, refused on submit
        "dropdown": {"option_id": o["dropdown"][1]},
        "image_choice": {"option_ids": o["image_choice"][:2]},
        "short_text": {"text": "not-an-email"},
        "text": {"text": "short"},
        "scale": {"rating": 0},
        "stars": {"rating": 7},
        "nps": {"rating": 9},
        "yes_no": {"unsure": True},
        "date": {"date": "2026-05-05"},
        "time": {"time": "09:30"},
        "number": {"number": 2.5},
        "matrix": {"matrix": {rows[0]: o["matrix"][0]}},
        "matrix_multi": {"matrix": {"r1": o["matrix_multi"][:2]}},
        "ranking": {"order": list(reversed(o["ranking"]))},
        "points": {"points": {str(o["points"][0]): 60, str(o["points"][1]): 30}},
        "pairwise": {"pairs": {f"{o['pairwise'][0]}-{o['pairwise'][1]}": o["pairwise"][1]}},
        "file": {"files": [uploaded]},
        "video": {"text": "Foydali"},
        "contact": {"contact": {"full_name": "Ali Valiyev"}},
        "region": {"region": tash, "district": None},
        "rating5": {"rating": 4},
    }
    for name, value in good.items():
        r = put(name, value)
        assert r.status_code == 200, (name, r.text)
    assert put("info", {"text": "x"}).json()["value"] is None

    r = client.post(f"/api/responses/{token}/submit", headers=USER)
    assert r.status_code == 422
    reasons = {m["question_id"]: m["reason"] for m in r.json()["detail"]["missing"]}
    by_id = {v["id"]: k for k, v in ids.items()}
    assert {by_id[k]: v for k, v in reasons.items()} == {
        "multiple": "min_choices",
        "short_text": "format",
        "text": "min_len",
        "number": "integer",
        "matrix": "incomplete",
        "points": "points_sum",
        "pairwise": "incomplete",
        "contact": "incomplete",
        "region": "incomplete",
    }

    fixes = {
        "multiple": {"option_ids": o["multiple"][:2]},
        "short_text": {"text": "ali@fti.uz"},
        "text": {"text": "Juda yaxshi xizmat"},
        "number": {"number": 25},
        "matrix": {"matrix": {rows[0]: o["matrix"][0], rows[1]: o["matrix"][2]}},
        "points": {"points": {str(o["points"][0]): 60, str(o["points"][1]): 30, str(o["points"][2]): 10}},
        "pairwise": {"pairs": {f"{a}-{b}": b for i, a in enumerate(sorted(o["pairwise"])) for b in sorted(o["pairwise"])[i + 1 :]}},
        "contact": {"contact": {"full_name": "Ali Valiyev", "phone": "+998 90 123 45 67"}},
        "region": {"region": tash, "district": chil},
    }
    for name, value in fixes.items():
        assert put(name, value).status_code == 200, name
    r = client.post(f"/api/responses/{token}/submit", headers=USER)
    assert r.status_code == 200, r.text

    # results
    res = client.get(f"/api/admin/surveys/{sid}/results", headers=ADMIN).json()
    stats = {by_id[x["id"]]: x["stats"] for x in res["questions"]}
    assert stats["nps"]["nps"]["score"] == 100
    assert stats["stars"]["kind"] == "rating" and len(stats["stars"]["items"]) == 7
    assert stats["matrix"]["kind"] == "matrix" and stats["matrix"]["rows"][1]["cells"][2]["count"] == 1
    assert stats["ranking"]["items"][0]["option_id"] == o["ranking"][-1]
    assert stats["points"]["items"][0]["count"] == 60
    assert stats["pairwise"]["items"][0]["option_id"] == max(o["pairwise"])
    assert stats["region"]["items"][0]["text_ru"] == "г. Ташкент"
    assert stats["region"]["districts"][0]["text_uz"] == "Chilonzor"
    assert stats["yes_no"]["items"][2]["count"] == 1
    assert stats["info"]["kind"] == "none"

    # response viewer and the file download link
    rid = client.get(f"/api/admin/surveys/{sid}/responses?status=completed", headers=ADMIN).json()["items"][0]["id"]
    detail = client.get(f"/api/admin/responses/{rid}?lang=ru", headers=ADMIN).json()
    answers = {it["question_id"]: it for it in detail["items"]}
    assert answers[ids["region"]["id"]]["answer"] == "г. Ташкент, Чиланзар"
    assert answers[ids["number"]["id"]]["answer"] == "25 лет"
    assert answers[ids["yes_no"]["id"]]["answer"] == "Затрудняюсь ответить"
    assert ids["info"]["id"] not in answers
    files = answers[ids["file"]["id"]]["files"]
    link = client.post(f"/api/admin/responses/{rid}/file-link", params={"key": files[0]["key"]}, headers=ADMIN).json()
    dl = client.get(link["url"])
    assert dl.status_code == 200 and dl.content == b"%PDF-1.4 test"
    assert dl.headers["content-type"] == "application/pdf"
    assert client.get(link["url"].replace("/answer-file/", "/answer-file/1")).status_code == 403
    # a respondent can't get a download link
    assert client.post(f"/api/admin/responses/{rid}/file-link", params={"key": files[0]["key"]}, headers=USER).status_code == 403

    # Excel export contains the readable answers
    x = client.get(f"/api/admin/surveys/{sid}/export.xlsx?lang=ru", headers=ADMIN)
    assert x.status_code == 200
    wb = load_workbook(BytesIO(x.content))
    values = [c for row in wb.worksheets[1].iter_rows(values_only=True) for c in row]
    assert "ali@fti.uz" in values and "г. Ташкент, Чиланзар" in values and "cv.pdf" in values


def test_settings_validation(client):
    sid = make_survey(client)
    url = f"/api/admin/surveys/{sid}/questions"
    bad = [
        q("scale", {"min": 1, "max": 200}),
        q("scale", {"min": 1, "max": 10, "step": 4}),
        q("stars", {"count": 20}),
        q("short_text", {"format": "passport"}),
        q("text", {"min_len": 50, "max_len": 10}),
        q("date", {"no_past": True, "no_future": True}),
        q("time", {"min_time": "9:00"}),
        q("number", {"min": 10, "max": 1}),
        q("matrix", {"rows": []}, opts()),
        q("file", {"formats": ["exe"]}),
        q("file", {"max_mb": 500}),
        q("contact", {"fields": {}}),
        q("pairwise", options=opts(7)),
        q("ranking", options=opts(1)),
        q("video", video_url="javascript:alert(1)"),
    ]
    for body in bad:
        assert client.post(url, json=body, headers=ADMIN).status_code == 422, body
    r = client.post(url, json=q("contact", {"fields": {"phone": True}}), headers=ADMIN)
    assert r.json()["detail"]["code"] == "contact_purpose_required"
    r = client.post(url, json=q("matrix", {"rows": [{"text_uz": "a", "text_ru": ""}]}, opts()), headers=ADMIN)
    assert r.json()["detail"]["code"] == "row_text_required"
    # unknown keys are dropped
    r = client.post(url, json=q("stars", {"count": 4, "evil": "<script>"}), headers=ADMIN)
    assert r.json()["settings"] == {"count": 4}


def test_contact_not_allowed_in_anonymous_survey(client):
    sid = make_survey(client, "anonymous")
    client.post(f"/api/admin/surveys/{sid}/questions", json=QUESTIONS["contact"], headers=ADMIN)
    r = client.post(f"/api/admin/surveys/{sid}/publish", headers=ADMIN)
    assert r.status_code == 422 and r.json()["detail"]["code"] == "personal_in_anonymous"


def test_only_info_blocks_cannot_be_published(client):
    sid = make_survey(client)
    client.post(f"/api/admin/surveys/{sid}/questions", json=QUESTIONS["info"], headers=ADMIN)
    assert client.post(f"/api/admin/surveys/{sid}/publish", headers=ADMIN).json()["detail"]["code"] == "no_questions"


def test_regions_admin(client):
    tash, sam, chil = seed_regions()
    admin = auth(7100, "Adm")
    client.get("/api/me", headers=admin)
    client.post("/api/admin/admins", json={"telegram_id": 7100, "role": "admin"}, headers=ADMIN)
    assert client.post("/api/admin/regions", json={"name_uz": "X", "name_ru": "X"}, headers=admin).status_code == 403

    r = client.post("/api/admin/regions", json={"name_uz": "Urgut", "name_ru": "Ургут", "parent_id": sam}, headers=ADMIN)
    assert r.status_code == 200
    new_id = r.json()["id"]
    assert client.post("/api/admin/regions", json={"name_uz": "Y", "name_ru": "Y", "parent_id": new_id}, headers=ADMIN).status_code == 422
    client.patch(f"/api/admin/regions/{new_id}", json={"name_ru": "Ургутский район"}, headers=ADMIN)
    tree = client.get("/api/admin/regions", headers=ADMIN).json()
    assert tree[1]["districts"][0]["name_ru"] == "Ургутский район"
    client.delete(f"/api/admin/regions/{tash}", headers=ADMIN)
    tree = client.get("/api/regions", headers=USER).json()
    assert [x["id"] for x in tree] == [sam]


def test_option_pictures_and_video_need_media_permission(client):
    sid = make_survey(client)
    editor = auth(7200, "Ed")
    client.get("/api/me", headers=editor)
    perms = {p: True for p in ("survey_create", "survey_edit", "question_manage", "question_type", "option_manage", "results_view")}
    client.post("/api/admin/admins", json={"telegram_id": 7200, "role": "admin", "permissions": perms, "surveys": [{"id": sid, "level": "edit"}]}, headers=ADMIN)
    url = f"/api/admin/surveys/{sid}/questions"
    assert client.post(url, json=q("image_choice", {}, opts(2, image=True)), headers=editor).status_code == 403
    assert client.post(url, json=q("video", video_url="https://youtu.be/x"), headers=editor).status_code == 403
    r = client.post(url, json=q("matrix", {"rows": [{"text_uz": "a", "text_ru": "а"}]}, opts(2)), headers=editor)
    assert r.status_code == 200
    # without option_manage the matrix rows can be renamed but not added
    perms["option_manage"] = False
    me = client.get("/api/admin/admins", headers=ADMIN).json()["items"]
    aid = next(a["id"] for a in me if a["telegram_id"] == 7200)
    client.patch(f"/api/admin/admins/{aid}", json={"permissions": perms}, headers=ADMIN)
    body = r.json()
    rows = body["settings"]["rows"]
    renamed = {**body, "settings": {"rows": [{**rows[0], "text_ru": "б"}]}}
    assert client.put(f"/api/admin/questions/{body['id']}", json=renamed, headers=editor).status_code == 200
    added = {**body, "settings": {"rows": rows + [{"text_uz": "c", "text_ru": "в"}]}}
    assert client.put(f"/api/admin/questions/{body['id']}", json=added, headers=editor).status_code == 403
